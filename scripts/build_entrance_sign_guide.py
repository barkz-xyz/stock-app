from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output" / "cartel-entrada"
OUTPUT_DOCX = OUTPUT_DIR / "guia-produccion-carteles-entrada.docx"

NAVY = "1C2630"
GOLD = "D98E00"
CREAM = "F6F0E7"
MUTED = "5F6670"
BLACK = "111827"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margins(cell, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run(run, size: float, color: str = BLACK, bold: bool = False, font: str = "Arial") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def set_paragraph(paragraph, before: float = 0, after: float = 6, line: float = 1.1) -> None:
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    set_paragraph(paragraph, before=12 if level == 1 else 8, after=6)
    run = paragraph.add_run(text)
    set_run(run, 16 if level == 1 else 13, NAVY if level == 1 else GOLD, True)


def add_body(doc: Document, text: str, bold_lead: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    set_paragraph(paragraph)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run(lead, 10.5, BLACK, True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run(rest, 10.5, BLACK)
    else:
        run = paragraph.add_run(text)
        set_run(run, 10.5, BLACK)


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    set_paragraph(paragraph, after=4, line=1.1)
    run = paragraph.add_run(text)
    set_run(run, 10.25, BLACK)


def set_table_geometry(table, widths_in: list[float], indent_dxa: int = 120) -> None:
    total = int(sum(widths_in) * 1440)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_in:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(width * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)


def build() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)

    for style_name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 13, GOLD)):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = hp.add_run("ATELIER DEL PUERTO  |  GUÍA DE PRODUCCIÓN")
    set_run(hr, 8, GOLD, True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Cartel de entrada · Archivos maestros incluidos en la misma carpeta")
    set_run(fr, 8, MUTED)

    kicker = doc.add_paragraph()
    set_paragraph(kicker, after=2)
    kr = kicker.add_run("PASTELERÍA ARTESANAL")
    set_run(kr, 9, GOLD, True)

    title = doc.add_paragraph()
    set_paragraph(title, after=4)
    tr = title.add_run("Guía de producción del cartel de entrada")
    set_run(tr, 24, NAVY, True)

    subtitle = doc.add_paragraph()
    set_paragraph(subtitle, after=14)
    sr = subtitle.add_run("Atelier del Puerto · Especificaciones para impresión, revisión y montaje")
    set_run(sr, 11.5, MUTED)

    summary = doc.add_table(rows=1, cols=3)
    summary.style = "Table Grid"
    summary_data = [
        ("HORIZONTAL", "142 × 64 cm", "Formato principal"),
        ("VERTICAL", "64 × 142 cm", "Formato alternativo"),
        ("PALETA", "Azul + dorado", "Alineada con la fachada"),
    ]
    for cell, (label, value, note) in zip(summary.rows[0].cells, summary_data):
        set_cell_shading(cell, NAVY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_paragraph(paragraph, after=2)
        r1 = paragraph.add_run(label + "\n")
        set_run(r1, 8.5, GOLD, True)
        r2 = paragraph.add_run(value + "\n")
        set_run(r2, 13, "FFFFFF", True)
        r3 = paragraph.add_run(note)
        set_run(r3, 8.5, "D7DDE3")
    set_table_geometry(summary, [2.17, 2.17, 2.16])

    add_heading(doc, "1. Archivos incluidos")
    files_table = doc.add_table(rows=1, cols=2)
    files_table.style = "Table Grid"
    headers = files_table.rows[0].cells
    headers[0].text = "Archivo"
    headers[1].text = "Uso"
    files = [
        ("cartel-entrada-horizontal-142x64cm.pdf", "Arte final horizontal a escala real, una página."),
        ("cartel-entrada-vertical-64x142cm.pdf", "Arte final vertical a escala real, una página."),
        ("cartel-entrada-horizontal-preview.png", "Vista previa para aprobar composición y contenido."),
        ("cartel-entrada-vertical-preview.png", "Vista previa para aprobar composición y contenido."),
    ]
    for file_name, use in files:
        cells = files_table.add_row().cells
        cells[0].text = file_name
        cells[1].text = use
    for cell in headers:
        set_cell_shading(cell, GOLD)
    for row_index, row in enumerate(files_table.rows):
        for col_index, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                set_paragraph(paragraph, after=0, line=1.0)
                for run in paragraph.runs:
                    set_run(run, 9.2, NAVY if row_index == 0 else BLACK, row_index == 0)
            if row_index > 0 and row_index % 2 == 0:
                set_cell_shading(cell, CREAM)
    set_table_geometry(files_table, [3.2, 3.3])

    add_heading(doc, "2. Contenido y enfoque visual")
    add_body(
        doc,
        "La propuesta retoma el azul profundo y las líneas doradas presentes en la fachada. "
        "Se utiliza un fondo oscuro de alto contraste, fotografía propia del negocio y bloques claros "
        "para que productos y precios puedan leerse desde la entrada.",
    )
    add_bullet(doc, "Selección compacta de 15 productos basada en el croquis: empanadas, panes/aliados, pizza y favoritos.")
    add_bullet(doc, "Productos ordenados de menor a mayor precio dentro de cada grupo.")
    add_bullet(doc, "Aviso mayorista visible: desde 12 unidades, precio a convenir por interno.")
    add_bullet(doc, "Fotografías tomadas exclusivamente del contenido disponible en la carpeta Imagenes.")

    add_heading(doc, "3. Recomendaciones de impresión")
    add_bullet(doc, "Imprimir los PDF al 100 % de escala. No usar opciones como Ajustar, Encoger o Fit to page.")
    add_bullet(doc, "Confirmar con la imprenta que el tamaño final sea exactamente 142 × 64 cm o 64 × 142 cm, según la orientación.")
    add_bullet(doc, "Material sugerido para exterior: PVC espumado de 5 mm o panel compuesto de aluminio de 3 mm con laminado UV mate.")
    add_bullet(doc, "Si se imprime en adhesivo, solicitar vinilo polimérico con laminado UV mate sobre una base rígida y plana.")
    add_bullet(doc, "Evitar acabado muy brillante: la fachada y la vitrina pueden generar reflejos que reducen la lectura.")
    add_bullet(doc, "Pedir una prueba de color pequeña para comprobar que el azul no salga morado y que el dorado conserve intensidad.")

    add_heading(doc, "4. Montaje y seguridad")
    add_bullet(doc, "Mantener al menos 2 cm libres desde el borde exterior al realizar perforaciones o instalar fijaciones.")
    add_bullet(doc, "Usar fijación mecánica acorde al muro y al peso final; evitar depender solo de cinta si estará expuesto a viento o humedad.")
    add_bullet(doc, "Colocar el horizontal con su lado largo paralelo al piso. El vertical debe instalarse con 142 cm de altura.")
    add_bullet(doc, "Revisar nivel, distancia de lectura y ausencia de reflejos antes de apretar definitivamente las fijaciones.")

    add_heading(doc, "5. Control antes de aprobar")
    checks = [
        "Precios y nombres revisados contra el catálogo vigente.",
        "Logo nítido y sin fondo blanco.",
        "Fotografías sin estiramiento ni recortes accidentales.",
        "Líneas doradas continuas y márgenes simétricos.",
        "Texto legible desde la distancia habitual de aproximación a la entrada.",
        "Orientación y medidas confirmadas por escrito con la imprenta.",
    ]
    for check in checks:
        paragraph = doc.add_paragraph()
        set_paragraph(paragraph, after=4)
        marker = paragraph.add_run("☐  ")
        set_run(marker, 11, GOLD, True, "Arial")
        run = paragraph.add_run(check)
        set_run(run, 10.25, BLACK)

    add_heading(doc, "Notas para la imprenta")
    note = doc.add_table(rows=1, cols=1)
    note.style = "Table Grid"
    cell = note.cell(0, 0)
    set_cell_shading(cell, CREAM)
    text = (
        "Los PDF fueron creados a tamaño final. Si el flujo de producción requiere sangrado, marcas de corte o conversión de color, "
        "debe agregarlos la imprenta sin alterar la proporción del arte. Solicitar una prueba antes de producir el panel definitivo."
    )
    cell.paragraphs[0].text = text
    for run in cell.paragraphs[0].runs:
        set_run(run, 9.8, NAVY, False)
    set_table_geometry(note, [6.5])

    doc.core_properties.title = "Guía de producción de carteles - Atelier del Puerto"
    doc.core_properties.subject = "Especificaciones para impresión y montaje"
    doc.core_properties.author = "Atelier del Puerto"
    doc.save(OUTPUT_DOCX)
    return OUTPUT_DOCX


if __name__ == "__main__":
    print(build())
